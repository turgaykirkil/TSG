#!/usr/bin/env python3
import os
import sys
from pathlib import Path
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

# Colors
DARK_BG = RGBColor(6, 11, 19)        # #060b13
CARD_BG = RGBColor(10, 22, 40)       # #0a1628
CARD_BORDER = RGBColor(56, 189, 248) # #38bdf8
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(148, 163, 184)     # #94a3b8
CYAN = RGBColor(56, 189, 248)       # #38bdf8
VIOLET = RGBColor(168, 85, 247)     # #a855f7
GREEN = RGBColor(34, 197, 94)       # #22c55e
RED = RGBColor(239, 68, 68)         # #ef4444

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_native_transition(slide):
    # Adds smooth PowerPoint native fade transition XML
    transition_xml = parse_xml(r"""
      <p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow">
        <p:fade/>
      </p:transition>
    """)
    slide._element.append(transition_xml)

def add_header(slide, badge_text, title_text, sub_text=""):
    # Badge
    badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf = badge_box.text_frame
    p = tf.paragraphs[0]
    p.text = badge_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf2 = title_box.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    
    # Subtitle
    if sub_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.6))
        tf3 = sub_box.text_frame
        p3 = tf3.paragraphs[0]
        p3.text = sub_text
        p3.font.size = Pt(13)
        p3.font.color.rgb = MUTED

def add_glass_card(slide, left, top, width, height, title, desc, icon="", title_color=WHITE, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = CARD_BG
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    p = tf.paragraphs[0]
    p.text = f"{icon} {title}".strip()
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = title_color
    
    if desc:
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED
        p2.space_before = Pt(6)

def build_presentation():
    print("🚀 Generating Native 100% Editable PowerPoint Presentation (PPTX)...")

    project_root = Path(__file__).parent.parent
    output_path = project_root / "sicilius" / "docs" / "Sicilius_Girisim_Sunumu_PitchDeck.pptx"
    public_path = project_root / "sicilius" / "frontend" / "public" / "presentation" / "Sicilius_Girisim_Sunumu_PitchDeck.pptx"
    
    output_path.parent.mkdir(parents=True, exist_ok=True)
    public_path.parent.mkdir(parents=True, exist_ok=True)

    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333) # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    assets_web = project_root / "sicilius" / "frontend" / "public" / "assets" / "web"
    assets_mobile = project_root / "sicilius" / "frontend" / "public" / "assets" / "mobile"

    # ==========================================
    # SLIDE 1: COVER HERO
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, DARK_BG)
    add_native_transition(s1)

    # Live Badge
    badge_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(0.5))
    p = badge_box.text_frame.paragraphs[0]
    p.text = "● CANLI PLATFORM — SICILIUS.COM.TR"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = CYAN

    # Main Title
    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.3), Inches(1.5))
    p = t_box.text_frame.paragraphs[0]
    p.text = "Türkiye'nin Akıllı Şirket\nİstihbarat Altyapısı"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Subtitle
    sub_box = s1.shapes.add_textbox(Inches(1.0), Inches(3.4), Inches(11.3), Inches(1.0))
    p = sub_box.text_frame.paragraphs[0]
    p.text = "Kamuya açık kaynaklardaki 2.100.000+ aktif şirket verisini ve 16.8M+ gazete ilanı sayfasını yapay zeka ile analiz eden, gizli ilişki ağlarını haritaya döken ve kurumsal kararları hızlandıran yenilikçi platform."
    p.font.size = Pt(15)
    p.font.color.rgb = MUTED

    # 4 Stat Cards
    stats_data = [
        ("2.1M+", "Aktif Şirket Kaydı"),
        ("16.8M+", "İlan & PDF Sayfası"),
        ("<200ms", "Arama Yanıtı"),
        ("100%", "Canlı & Otonom")
    ]
    for idx, (num, label) in enumerate(stats_data):
        left = Inches(1.0 + idx * 2.9)
        top = Inches(4.8)
        shape = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(2.7), Inches(1.6))
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = CARD_BORDER
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        p1 = tf.paragraphs[0]
        p1.text = num
        p1.font.size = Pt(28)
        p1.font.bold = True
        p1.font.color.rgb = CYAN
        p1.alignment = PP_ALIGN.CENTER

        p2 = tf.add_paragraph()
        p2.text = label
        p2.font.size = Pt(12)
        p2.font.color.rgb = MUTED
        p2.alignment = PP_ALIGN.CENTER
        p2.space_before = Pt(4)

    # ==========================================
    # SLIDE 2: PROBLEM
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, DARK_BG)
    add_native_transition(s2)
    add_header(s2, "01 · PROBLEM", "Mevcut Şirket Araştırma Sistemleri Neden Yetersiz?", "Kamuya açık kaynaklardaki şirket verileri erişilebilir olmasına rağmen aranması ve ilişkilendirilmesi son derece zordur.")

    problems = [
        ("📰", "PDF Dağınıklığı", "Kamuya açık ilanlar yüzlerce sayfalık araması imkansız ham PDF belgeleri halinde yayınlanmaktadır."),
        ("🕸️", "Görünmeyen Ağlar", "Bir şirketin ortak yöneticileri, gizli bağları ve yan iştiraklerini tespit etmek saatler alır."),
        ("⏳", "Zaman & İş Gücü Kaybı", "Kredi risk ve KYC araştırmaları manuel taramalar nedeniyle günlerce sürmektedir."),
        ("💰", "Yüksek Maliyetler", "Piyasadaki ham veri satıcıları yüksek lisans ücretleri talep etmekte ve hantal kalmaktadır.")
    ]
    for idx, (icon, title, desc) in enumerate(problems):
        left = Inches(0.8 + idx * 2.95)
        top = Inches(2.4)
        add_glass_card(s2, left, top, Inches(2.8), Inches(4.2), title, desc, icon=icon, title_color=WHITE, border_color=RED)

    # ==========================================
    # SLIDE 3: ARAMA DENEYİMİ (WITH IMAGE)
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, DARK_BG)
    add_native_transition(s3)
    add_header(s3, "02 · ARAMA DENEYİMİ", "Minimalist & Güçlü Arama Deneyimi", "Karmaşık kamu verisini Google sadeliğinde tek bir arama çubuğuna indirgeyen sezgisel arayüz.")

    add_glass_card(s3, Inches(0.8), Inches(2.2), Inches(5.4), Inches(1.4), "Google Sadeliğinde Arama", "Tek bir alana VKN, Şirket Unvanı veya MERSİS No yazarak anında sorgulama yapın.", icon="🔍")
    add_glass_card(s3, Inches(0.8), Inches(3.8), Inches(5.4), Inches(1.4), "Hızlı Kısayol Tuşları", "CMD + K veya / kısayolları ile platformun her yerinden saniyede yeni arama başlatın.", icon="⚡")
    add_glass_card(s3, Inches(0.8), Inches(5.4), Inches(5.4), Inches(1.4), "Güvenli & Standart Veri", "Türkçe karakter normalizasyonu ve fuzzy match motoru ile hatalı yazımlarda bile anında bulun.", icon="🛡️")

    img_path_search = assets_web / "web_search_dashboard.png"
    if img_path_search.exists():
        s3.shapes.add_picture(str(img_path_search), Inches(6.5), Inches(2.2), width=Inches(6.0))

    # ==========================================
    # SLIDE 4: ÇÖZÜM & NEXUS (WITH IMAGE)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, DARK_BG)
    add_native_transition(s4)
    add_header(s4, "03 · ÇÖZÜM & NEXUS", "Sicilius ile Akıllı Kurumsal İstihbarat", "Ham kamu verisini saniyeler içinde karar verilebilir görsel analize dönüştürüyoruz.")

    img_path_nexus = assets_web / "web_nexus_graph.png"
    if img_path_nexus.exists():
        s4.shapes.add_picture(str(img_path_nexus), Inches(0.8), Inches(2.2), width=Inches(6.5))

    add_glass_card(s4, Inches(7.6), Inches(2.2), Inches(4.9), Inches(1.4), "AI Destekli OCR + NLP Engine", "Kamuya açık ilanları otonom tarar, optik karakter tanıma ile metne çevirir.", icon="🤖")
    add_glass_card(s4, Inches(7.6), Inches(3.8), Inches(4.9), Inches(1.4), "NEXUS İlişki Ağ Grafiği", "Kişi-şirket bağlantılarını Depth-2 derinliğinde görselleştirir, anomali tespiti yapar.", icon="🕸️")
    add_glass_card(s4, Inches(7.6), Inches(5.4), Inches(4.9), Inches(1.4), "<200ms Ultra Hızlı Arama", "In-memory TTL Önbellekleme ve PostGIS spatial veritabanı indekslemesi ile anlık yanıtlar.", icon="⚡")

    # ==========================================
    # SLIDE 5: YETENEKLER
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, DARK_BG)
    add_native_transition(s5)
    add_header(s5, "04 · YETENEKLER", "Uçtan Uca Platform Modülleri", "Bir kurumsal araştırma ve istihbarat merkezinin ihtiyaç duyduğu tüm bileşenler.")

    tiles = [
        ("🔍", "Akıllı Şirket Arama", "Fuzzy match, Türkçe karakter normalizasyonu ve VKN / Sicil No ile anlık arama."),
        ("🕸️", "NEXUS Grafik Motoru", "Şirket-Kişi graf haritası, ML anomali tespiti ve dolaylı ortaklık araştırması."),
        ("📄", "OCR + İlan İşleme", "Kamuya açık PDF belgelerini otonom tarar, NLP ile yönetici değişikliklerini ayıklar."),
        ("🗺️", "Coğrafi Haritalama", "PostGIS altyapısı, otomatik geocoding ve interaktif kümeleme haritası."),
        ("🏆", "Oyunlaştırma & Saha", "GPS 50m doğrulama ile saha check-in, anonim KVKK liderlik tablosu ve XP sistemi."),
        ("🐹", "Go (Golang) Microservice", "PostGIS ST_DWithin ile 50m saha fırsatları sunan Clean Architecture REST API.")
    ]
    for idx, (icon, title, desc) in enumerate(tiles):
        row = idx // 3
        col = idx % 3
        left = Inches(0.8 + col * 3.95)
        top = Inches(2.2 + row * 2.4)
        add_glass_card(s5, left, top, Inches(3.8), Inches(2.2), title, desc, icon=icon)

    # ==========================================
    # SLIDE 6: TEKNOLOJİ & MİMARİ
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, DARK_BG)
    add_native_transition(s6)
    add_header(s6, "05 · TEKNOLOJİ", "Modern & Kurumsal Mimari Altyapı", "Yüksek performanslı, containerized ve sürekli çalışan mikro-servis mimarisi.")

    techs = [
        ("FRONTEND", "Next.js 14 & React", "TypeScript, Tailwind CSS, React Force Graph entegrasyonu."),
        ("BACKEND", "Go (Golang) + FastAPI", "Clean Architecture, Python 3.11 NLP işleme ve high-concurrency API."),
        ("VERİTABANI", "PostgreSQL + PostGIS", "Coğrafi indeksleme, MinIO S3 dosya depolama ve in-memory cache."),
        ("MOBİL", "Expo (React Native)", "OpenStreetMap ücretsiz harita altyapısı, Caddy TLS ve DR planı.")
    ]
    for idx, (cat, title, desc) in enumerate(techs):
        row = idx // 2
        col = idx % 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(2.2 + row * 2.4)
        add_glass_card(s6, left, top, Inches(5.6), Inches(2.2), f"[{cat}] {title}", desc, icon="⚡")

    # ==========================================
    # SLIDE 7: SAHA SATIŞI & GAMIFICATION (WITH 2 IPHONE IMAGES)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, DARK_BG)
    add_native_transition(s7)
    add_header(s7, "06 · SAHA SATIŞI & GİZLİLİK", "Saha Ziyaret Disiplini & Oyunlaştırma Engine", "KVKK uyumlu anonim sıralama ve GPS 50m doğrulama ile saha ekiplerini motive edin.")

    add_glass_card(s7, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.8), "5 Kademeli Lig & Anonim Tablo", "Bronz -> Gümüş -> Altın -> Platin -> Efsane Lig sistemi ve KVKK uyumlu anonim liderlik tablosu.", icon="🏆")
    img_path_gami = assets_mobile / "mobile_gamification.png"
    if img_path_gami.exists():
        s7.shapes.add_picture(str(img_path_gami), Inches(3.6), Inches(3.0), height=Inches(3.8))

    add_glass_card(s7, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.8), "%100 Ücretsiz Harita Altyapısı", "OpenStreetMap (OSM) Tile Server entegrasyonu ile sıfır API maliyetli harita renderlama.", icon="🗺️")
    img_path_osm = assets_mobile / "mobile_map_osm.png"
    if img_path_osm.exists():
        s7.shapes.add_picture(str(img_path_osm), Inches(9.6), Inches(3.0), height=Inches(3.8))

    # ==========================================
    # SLIDE 8: PAZAR & KULLANIM ALANLARI
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, DARK_BG)
    add_native_transition(s8)
    add_header(s8, "07 · PAZAR & KULLANIM ALANLARI", "Pazar Fırsatı ve Kullanım Senaryoları", "Türkiye'de kurumsal veri ve regtech alanındaki yüksek büyüme potansiyeli.")

    uses = [
        ("🏦", "Bankacılık & Finans", "KYC, kredi risk değerlendirmesi ve müşteri due diligence süreçlerinde hızlı şirket ve yönetici geçmişi araştırması."),
        ("⚖️", "Hukuk & Danışmanlık", "M&A ortaklık süreçleri, dava araştırmaları, iştirak ağ haritalandırması ve hukuki incelemeler."),
        ("🏢", "Saha Satış & KOBİ", "Yeni kurulan şirketlerin tespiti, coğrafi yakınlık araması ve saha ziyaret disiplini optimizasyonu.")
    ]
    for idx, (icon, title, desc) in enumerate(uses):
        left = Inches(0.8 + idx * 3.95)
        top = Inches(2.4)
        add_glass_card(s8, left, top, Inches(3.8), Inches(4.2), title, desc, icon=icon)

    # ==========================================
    # SLIDE 9: YOL HARİTASI
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, DARK_BG)
    add_native_transition(s9)
    add_header(s9, "08 · YOL HARİTASI", "Gelişim ve Gelecek Vizyonu", "Tamamlanan kilometre taşları ve önümüzdeki dönemin teknolojik hedefleri.")

    phases = [
        ("✅ 2025 Q1-Q2", "Temel Altyapı", "FastAPI + Next.js Kurulumu, PostgreSQL + PostGIS Coğrafi Altyapı, Otonom Data Pipeline"),
        ("✅ 2025 Q3-Q4", "NEXUS & Mobil", "NEXUS Derinlik-2 Graf Motoru, React Native Mobil App, Go Microservice & DR Planı"),
        ("⚡ 2026 Q1", "Saha & Oyunlaştırma", "5 Kademeli Lig, GPS 50m Doğrulama, Anonim Tablo & 1.000+ Dinamik Görev Motoru"),
        ("🚀 2026 Q2+", "AI & Gelecek", "LLM Destekli Özet Raporlama, Yapay Zeka Şirket Karnesi & Otonom Tedarikçi Asistanı")
    ]
    for idx, (badge, title, desc) in enumerate(phases):
        left = Inches(0.8 + idx * 2.95)
        top = Inches(2.4)
        add_glass_card(s9, left, top, Inches(2.8), Inches(4.2), f"{badge}\n{title}", desc)

    # ==========================================
    # SLIDE 10: SSS
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, DARK_BG)
    add_native_transition(s10)
    add_header(s10, "09 · SSS", "Merak Edilen Sorular & Yanıtlar", "Sicilius altyapısı, veri kaynağı, güvenlik ve kullanım detayları hakkında merak edilenler.")

    faqs = [
        ("⚖️ Veri nereden geliyor, yasal altyapısı nedir?", "Kamuya açık resmi devlet yayınları ve açık kaynak veriler kullanılmaktadır. Sicilius ham kamu verisini yapay zeka ile yapılandırılmış ve sorgulanabilir hale getirir. KVKK standartlarına tam uyumludur."),
        ("🚀 Sicilius'u rakiplerinden ayıran en temel 3 fark nedir?", "(1) NEXUS ilişki ağı grafı (Depth-2), (2) GPS 50m doğrulama ve oyunlaştırmalı B2B saha satış modülü, (3) %100 otonom veri güncelleme ve sıfır harita maliyetli mimari."),
        ("🛡️ Oyunlaştırmada KVKK ve gizlilik nasıl sağlanıyor?", "Ciro veya satış tutarı kesinlikle takip edilmez. Liderlik tablosu tamamen anonimdir ('Temsilci #1', 'SEN 4. Sıra'). Diğer kullanıcıların kişisel bilgileri sistemde gizli tutulur."),
        ("🗺️ Mobil haritada Google Maps API ücreti ödeniyor mu?", "Hayır. Ücretli Google Maps API yerine %100 ücretsiz OpenStreetMap (OSM) Tile Server kullanılmıştır. Sıfır API maliyetiyle sınırsız harita render edilmektedir.")
    ]
    for idx, (q, a) in enumerate(faqs):
        left = Inches(0.8 + (idx % 2) * 5.9)
        top = Inches(2.2 + (idx // 2) * 2.4)
        add_glass_card(s10, left, top, Inches(5.6), Inches(2.2), q, a)

    # ==========================================
    # SLIDE 11: KAPANIS / CTA
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, DARK_BG)
    add_native_transition(s11)

    t_box = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(11.3), Inches(1.5))
    p = t_box.text_frame.paragraphs[0]
    p.text = "Türkiye'nin Şirket Verisini\nGeleceğe Taşıyoruz"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

    sub_box = s11.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(11.3), Inches(0.8))
    p = sub_box.text_frame.paragraphs[0]
    p.text = "Binlerce şirket ve yönetici ilişkisini saniyeler içinde keşfetmek için platformu hemen canlıda deneyimleyin."
    p.font.size = Pt(16)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.CENTER

    cta_box = s11.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.0))
    p = cta_box.text_frame.paragraphs[0]
    p.text = "🌐 sicilius.com.tr   |   📧 info@sicilius.com.tr   |   🔬 sicilius.com.tr/nexus"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN
    p.alignment = PP_ALIGN.CENTER

    prs.save(str(output_path))
    prs.save(str(public_path))

    print(f"🎉 Native PowerPoint (.pptx) generated successfully!")
    print(f"   📁 Output: {output_path} ({output_path.stat().st_size // 1024} KB)")
    print(f"   📁 Public: {public_path} ({public_path.stat().st_size // 1024} KB)")

if __name__ == "__main__":
    build_presentation()
