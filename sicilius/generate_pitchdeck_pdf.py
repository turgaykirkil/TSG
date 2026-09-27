#!/usr/bin/env python3
import os
import sys
import time
import subprocess
from pathlib import Path
import fitz  # PyMuPDF
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

def generate_pdf_and_pptx():
    print("🚀 Starting Automatic Pitch Deck PDF & PPTX Generation via Native Chrome...")

    # Paths
    project_root = Path(__file__).parent.parent
    frontend_dir = project_root / "sicilius" / "frontend"
    output_dir = project_root / "sicilius" / "docs"
    assets_dir = output_dir / "assets" / "presentation"
    public_presentation_dir = frontend_dir / "public" / "presentation"

    output_dir.mkdir(parents=True, exist_ok=True)
    assets_dir.mkdir(parents=True, exist_ok=True)
    public_presentation_dir.mkdir(parents=True, exist_ok=True)

    pdf_output_path = output_dir / "Sicilius_Girisim_Sunumu_PitchDeck.pdf"
    pptx_output_path = output_dir / "Sicilius_Girisim_Sunumu_PitchDeck.pptx"
    
    public_pdf_path = public_presentation_dir / "Sicilius_Girisim_Sunumu_PitchDeck.pdf"
    public_pptx_path = public_presentation_dir / "Sicilius_Girisim_Sunumu_PitchDeck.pptx"

    chrome_path = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
    if not os.path.exists(chrome_path):
        print(f"❌ Error: Chrome binary not found at {chrome_path}")
        sys.exit(1)

    url = "http://localhost:3000/presentation/index.html"
    print(f"🔗 Target presentation URL: {url}")

    tmp_dir = project_root / "sicilius" / "docs" / "assets" / "tmp_slides"
    tmp_dir.mkdir(parents=True, exist_ok=True)

    total_slides = 11
    slide_images = []

    print("📸 Capturing high-res 1920x1080 slide screenshots using Native Chrome Headless...")
    for slide_idx in range(1, total_slides + 1):
        png_filename = f"slide_{slide_idx:02d}.png"
        png_path = tmp_dir / png_filename
        presentation_asset_path = assets_dir / png_filename

        style_overrides = f"""
          * {{ transition: none !important; animation: none !important; }}
          .deck-header, .deck-controls, .overview-modal, .slide-footer-hint {{ display: none !important; }}
          .orb {{ opacity: 0.15 !important; }}
          .slide {{ 
            display: flex !important; 
            opacity: 1 !important; 
            visibility: visible !important;
            transform: none !important; 
            position: fixed !important; 
            inset: 0 !important; 
            padding: 50px 60px !important; 
            z-index: 100 !important;
            background: #060b13 !important;
          }}
          .slide:not([data-slide="{slide_idx}"]) {{ 
            display: none !important; 
            opacity: 0 !important; 
            visibility: hidden !important; 
          }}
          .slide[data-slide="{slide_idx}"] * {{ 
            opacity: 1 !important; 
            transform: none !important; 
          }}
        """

        cmd = [
          chrome_path,
          "--headless=new",
          "--disable-gpu",
          "--no-sandbox",
          "--window-size=1920,1080",
          f"--user-style-rules={style_overrides}",
          f"--screenshot={png_path}",
          url
        ]

        subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        time.sleep(0.5)

        if png_path.exists() and png_path.stat().st_size > 0:
            print(f"  ✅ Captured Slide {slide_idx}/{total_slides} -> {png_filename} ({png_path.stat().st_size // 1024} KB)")
            slide_images.append(png_path)
            # Copy to assets dir
            with open(png_path, 'rb') as src, open(presentation_asset_path, 'wb') as dst:
                dst.write(src.read())
        else:
            print(f"  ❌ Failed to capture Slide {slide_idx}")

    if not slide_images:
        print("❌ Error: No slide images were captured.")
        sys.exit(1)

    print(f"\n📦 Assembling 16:9 Landscape PDF Pitch Deck ({len(slide_images)} slides)...")

    # PyMuPDF High-Res PDF
    pdf_doc = fitz.open()
    width_pt = 1920 * 72 / 96  # 1440 pt
    height_pt = 1080 * 72 / 96 # 810 pt
    rect = fitz.Rect(0, 0, width_pt, height_pt)

    for img_path in slide_images:
        page = pdf_doc.new_page(width=width_pt, height=height_pt)
        page.insert_image(rect, filename=str(img_path))

    pdf_doc.save(str(pdf_output_path))
    pdf_doc.close()
    
    # Copy PDF to public presentation dir for instant web download
    with open(pdf_output_path, 'rb') as src, open(public_pdf_path, 'wb') as dst:
        dst.write(src.read())

    print(f"  🎉 PDF generated successfully! Path: {pdf_output_path} ({pdf_output_path.stat().st_size // 1024} KB)")

    # PowerPoint PPTX Widescreen (16:9)
    print(f"\n📊 Assembling 16:9 Widescreen PowerPoint Deck (PPTX)...")
    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 aspect ratio standard
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    for img_path in slide_images:
        slide = prs.slides.add_slide(blank_slide_layout)
        slide.shapes.add_picture(str(img_path), Inches(0), Inches(0), Inches(13.333), Inches(7.5))

    prs.save(str(pptx_output_path))
    
    # Copy PPTX to public presentation dir for instant web download
    with open(pptx_output_path, 'rb') as src, open(public_pptx_path, 'wb') as dst:
        dst.write(src.read())

    print(f"  🎉 PPTX generated successfully! Path: {pptx_output_path} ({pptx_output_path.stat().st_size // 1024} KB)")
    print("\n✨ Pitch deck PDF & PPTX export complete!")

if __name__ == "__main__":
    generate_pdf_and_pptx()
